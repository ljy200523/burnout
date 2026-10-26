from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR

OUT = r'C:\Users\USER\burnout\pace_reward_system.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Palette
BG = 'F7F7F2'; INK = '18332B'; MUTED = '64746A'; GREEN = '315D4D'
LIGHT = 'E8EFE6'; GOLD = 'C39A5A'; WHITE = 'FFFFFF'; LINE = 'DDE5DB'

def rgb(s): return RGBColor.from_string(s)
def shape(kind, x,y,w,h, fill, line=None, radius=None):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    sh.line.fill.background() if line is None else None
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(1)
    return sh
def textbox(x,y,w,h,text,size,color=INK,bold=False,font='Aptos',align=PP_ALIGN.LEFT,
            valign=MSO_ANCHOR.MIDDLE,margin=0):
    tb=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=tb.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin)
    tf.margin_top=Inches(0); tf.margin_bottom=Inches(0); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    return tb
def pill(x,y,w,label,fill=LIGHT,color=GREEN):
    sh=shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,.34,fill)
    sh.adjustments[0]=.18
    textbox(x+.1,y,w-.2,.34,label,10,color,True,align=PP_ALIGN.CENTER)

# Background and title
slide.background.fill.solid(); slide.background.fill.fore_color.rgb=rgb(BG)
pill(.66,.42,1.28,'PACE  /  REWARDS',fill='E9E5D8',color='776443')
textbox(.66,.92,12.05,.57,'Small rewards make consistency feel worthwhile.',30,INK,True)
textbox(.68,1.55,11.85,.36,'Check in, earn care points, and use them for simple everyday perks.',15,MUTED)

# Primary product panel
panel=shape(MSO_SHAPE.ROUNDED_RECTANGLE,.66,2.12,12.0,2.45,GREEN)
panel.adjustments[0]=.06
textbox(1.02,2.42,5.2,.26,'YOUR PACE, YOUR PROGRESS',10,'D9E4D7',True)
textbox(1.02,2.79,5.7,.48,'Keep your streak moving',23,WHITE,True)
textbox(1.02,3.34,5.1,.34,'Miss a day? A grace day protects your check-in streak.',13,'E6EEE7')

# Grace day tile
tile=shape(MSO_SHAPE.ROUNDED_RECTANGLE,7.03,2.47,2.32,1.56,'416B59')
tile.adjustments[0]=.08
textbox(7.25,2.67,1.88,.28,'GRACE DAY',10,'DCE9DD',True,align=PP_ALIGN.CENTER)
textbox(7.25,2.99,1.88,.42,'1 ready',23,WHITE,True,align=PP_ALIGN.CENTER)
textbox(7.25,3.48,1.88,.27,'Use when you miss a day',10,'E6EEE7',align=PP_ALIGN.CENTER)

# Points tile
tile=shape(MSO_SHAPE.ROUNDED_RECTANGLE,9.61,2.47,2.62,1.56,'416B59')
tile.adjustments[0]=.08
textbox(9.83,2.67,2.18,.28,'AVAILABLE BALANCE',10,'DCE9DD',True,align=PP_ALIGN.CENTER)
textbox(9.83,2.93,2.18,.57,'710',32,WHITE,True,align=PP_ALIGN.CENTER)
textbox(9.83,3.51,2.18,.25,'care points',11,'E6EEE7',align=PP_ALIGN.CENTER)

# Connector/callout labels
textbox(7.03,4.64,2.32,.23,'STREAK SAFETY',9,GREEN,True,align=PP_ALIGN.CENTER)
textbox(9.61,4.64,2.62,.23,'REWARD CURRENCY',9,GREEN,True,align=PP_ALIGN.CENTER)

# Rewards row heading
textbox(.68,5.02,8.6,.36,'Choose a small reward',20,INK,True)
textbox(9.3,5.07,3.35,.28,'Examples from the perks catalogue',11,MUTED,align=PP_ALIGN.RIGHT)

# Three compact cards
cards=[
    ("COFFEE / TEA","Campus café voucher","A coffee, matcha, or bubble tea treat.","100",'F0DFBF','7A5732','☕'),
    ("FOOD","Balanced canteen meal","A meal at a participating campus eatery.","180",'E7EBCB','4E693D','✦'),
    ("MOVEMENT / REST","Campus gym or pool pass","A guest pass or locker credit.","220",'DCE9DF','426B58','↗'),
]
xs=[.66,4.73,8.80]
for i,(cat,title,desc,cost,accent,accent_text,icon) in enumerate(cards):
    x=xs[i]
    c=shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,5.48,3.88,1.48,WHITE,LINE); c.adjustments[0]=.06
    # small accent strip
    strip=shape(MSO_SHAPE.ROUNDED_RECTANGLE,x+.17,5.70,.55,.55,accent); strip.adjustments[0]=.16
    textbox(x+.17,5.70,.55,.55,icon,19,accent_text,True,align=PP_ALIGN.CENTER)
    textbox(x+.84,5.68,2.86,.22,cat,8,accent_text,True)
    textbox(x+.84,5.94,2.86,.31,title,13,INK,True)
    textbox(x+.22,6.31,3.43,.28,desc,9,MUTED)
    textbox(x+.22,6.64,.84,.25,cost+' pts',12,GREEN,True)
    # compact redeem cue
    pill(x+2.82,6.62,.84,'REDEEM',fill=LIGHT,color=GREEN)

# Footer takeaway
textbox(.69,7.12,11.95,.18,'CHECK IN  →  EARN CARE POINTS  →  REDEEM A PERK',9,MUTED,True,align=PP_ALIGN.CENTER)

prs.core_properties.title='Pace — Reward System'
prs.core_properties.subject='Presentation slide: check-in streaks, grace days, and rewards'
prs.core_properties.author='OpenAI'
prs.save(OUT)
print(OUT)
